"""
Find dangling FHIR references (e.g. OrganizationAffiliation -> Location/4136336 that doesn't exist).

Usage:
    python check_hy.py H1619        # one contract (reads ./H1619/*.json)
    python check_hy.py              # all contracts

Reads each contract's bundle files straight from its local folder
(<script_dir>/<CONTRACT>/*.json) -- no network/index download involved, then:
  Pass 1 - build a set of every (resourceType, id) that actually exists.
  Pass 2 - walk every resource for "reference" fields and flag targets missing from that set.

Outputs (per contract):
  dangling_refs_<C>.csv      one row per broken reference
  dangling_summary_<C>.csv   affected counts + % per source type -> target type
"""
import json, csv, re, os, sys, glob
from datetime import date



# Prefix-agnostic: matches jhp-H1619-2027-location-part1.json AND any other org's
# <prefix>-<contract>-<year>-<category>-part<n>.json layout (e.g. CHPW's H5826 files).
FNAME_RE = re.compile(r"(?P<contract>[hH]\d+)-(?P<year>\d+)-(?P<category>[a-z]+)-part(?P<part>\d+)\.json", re.I)
# "Location/4136336", "Location/4136336/_history/1", or an absolute URL ending in Type/id
REF_RE = re.compile(r"(?:^|/)(?P<type>[A-Z][A-Za-z]+)/(?P<id>[A-Za-z0-9\-.]{1,64})(?:/_history/.*)?$")

PLACEHOLDER_EXT_URL = "http://hapifhir.io/fhir/StructureDefinition/resource-placeholder"


def has_placeholder_extension(resource):
    """True if the resource itself is explicitly marked placeholder data via the
    official 'resource-placeholder' extension (valueBoolean=true) -- the authoritative
    signal, as opposed to guessing from the shape of the id (see is_placeholder_id)."""
    for ext in resource.get("extension", []) or []:
        if ext.get("url") == PLACEHOLDER_EXT_URL and ext.get("valueBoolean") is True:
            return True
    return False

EXPECTED_ORPHAN_TYPES = {"payer", "network", "ntwk"}


def is_expected_orphan_type(type_str):
    low = (type_str or "").lower()
    return any(t in low for t in EXPECTED_ORPHAN_TYPES)


PLACEHOLDER_LITERALS = {
    "test", "example", "unknown", "tbd", "n/a", "na", "none", "null",
    "sample", "dummy", "fake", "todo", "xxx", "0000000000",
}


def is_placeholder_id(rid):
    """True if a target id looks like dummy/test data rather than a real FHIR id
    (all-zero, all-same-digit, simple sequential digits, or a known dummy literal),
    even though it may still happen to match a real resource in `existing`."""
    s = str(rid).strip()
    low = s.lower()
    if low in PLACEHOLDER_LITERALS:
        return True
    if s.isdigit():
        if len(set(s)) == 1:          # e.g. "0000", "1111"
            return True
        asc = "".join(str((int(s[0]) + i) % 10) for i in range(len(s)))
        desc = "".join(str((int(s[0]) - i) % 10) for i in range(len(s)))
        if len(s) >= 4 and s in (asc, desc):   # e.g. "1234", "9876"
            return True
    return False


def local_files_for(folder):
    """Every *.json file sitting in <folder>, sorted for stable output."""
    return sorted(glob.glob(os.path.join(folder, "*.json")))


def read_bytes(path):
    with open(path, "rb") as f:
        return f.read()


def resources_of(raw):
    """Yield each FHIR resource from a Bundle (or a bare list/object)."""
    data = json.loads(raw)
    if isinstance(data, dict) and data.get("resourceType") == "Bundle":
        for e in data.get("entry", []):
            r = e.get("resource", e)
            if isinstance(r, dict):
                yield r
    elif isinstance(data, list):
        for r in data:
            if isinstance(r, dict):
                yield r
    elif isinstance(data, dict):
        yield data


def find_refs(node, path=""):
    """Recursively yield (json_path, reference_string) for every 'reference' field."""
    if isinstance(node, dict):
        for k, v in node.items():
            p = f"{path}.{k}" if path else k
            if k == "reference" and isinstance(v, str):
                yield p, v
            else:
                yield from find_refs(v, p)
    elif isinstance(node, list):
        for i, v in enumerate(node):
            yield from find_refs(v, f"{path}[{i}]")


def field_root(path):
    """'location[0].reference' -> 'location'; 'coverageArea[3].reference' -> 'coverageArea'.
    Different fields mean different things (location vs coverageArea), so they are
    reported separately rather than collapsed into one source->target row."""
    p = re.sub(r"\[\d+\]", "", path)
    if p.endswith(".reference"):
        p = p[: -len(".reference")]
    return p or "reference"


def parse_ref(ref):
    """Return (type, id) for a resolvable reference, else None for contained/urn refs."""
    if not ref or ref.startswith("#") or ref.startswith("urn:"):
        return None
    m = REF_RE.search(ref.strip())
    if not m:
        return None
    return m.group("type"), m.group("id")


def run(contract, folder):
    print(f"\n=== {contract} ===", flush=True)
    paths = local_files_for(folder)
    if not paths:
        print(f"No *.json files found in {folder} -- skipping {contract}.", flush=True)
        return 0
    print(f"{len(paths)} local files in {folder}", flush=True)

    files = []  # (path, fname, category, part)
    for path in paths:
        fname = os.path.basename(path)
        m = FNAME_RE.search(fname)
        files.append((path, fname,
                      m.group("category") if m else "unknown",
                      int(m.group("part")) if m else 0))

    # ---- Pass 1: every (resourceType, id) that exists ----
    print("Pass 1: indexing existing resource ids ...", flush=True)
    existing = set()
    identifiers = {}              # (resourceType, id) -> "system|value" identifiers, joined by "; "
    org_meta = {}                  # (resourceType, id) -> {"name": ..., "type": "Payer; Network"}
    raw_type_counts = {}          # includes duplicate ids (raw entry count)
    ext_placeholder_by_file = {}  # fname -> count of resources with resource-placeholder ext
    for path, fname, category, part in files:
        raw = read_bytes(path)
        n = 0
        for r in resources_of(raw):
            rt, rid = r.get("resourceType"), r.get("id")
            if rt and rid is not None and str(rid) != "":
                existing.add((rt, str(rid)))
                raw_type_counts[rt] = raw_type_counts.get(rt, 0) + 1
                idents = [f"{i.get('system', '')}|{i.get('value', '')}"
                          for i in (r.get("identifier") or []) if isinstance(i, dict)]
                if idents:
                    identifiers[(rt, str(rid))] = "; ".join(idents)
                if rt in ("Organization", "Practitioner"):
                    type_labels = []
                    for t in (r.get("type") or []):
                        if not isinstance(t, dict):
                            continue
                        if t.get("text"):
                            type_labels.append(t["text"])
                        for coding in t.get("coding", []) or []:
                            if coding.get("display"):
                                type_labels.append(coding["display"])
                            elif coding.get("code"):
                                type_labels.append(coding["code"])
                    name = r.get("name")
                    if isinstance(name, list):  # Practitioner.name is HumanName[]
                        name = " ".join(
                            " ".join(n.get("given", []) + [n.get("family", "")])
                            for n in name if isinstance(n, dict)
                        ).strip()
                    org_meta[(rt, str(rid))] = {
                        "name": name or "",
                        "type": "; ".join(dict.fromkeys(type_labels)),
                    }
            if has_placeholder_extension(r):
                ext_placeholder_by_file[fname] = ext_placeholder_by_file.get(fname, 0) + 1
            n += 1
        print(f"  {fname}: {n}", flush=True)

    if ext_placeholder_by_file:
        print(f"\nResources marked with resource-placeholder extension, by file:")
        for fname in sorted(ext_placeholder_by_file):
            print(f"  {fname}: {ext_placeholder_by_file[fname]:,}")
        print(f"Total: {sum(ext_placeholder_by_file.values()):,}")
    # Unique count per type = correct denominator for "% of source affected".
    type_counts = {}
    for t, _ in existing:
        type_counts[t] = type_counts.get(t, 0) + 1
    published_types = set(type_counts)
    dups = sum(raw_type_counts[t] - type_counts.get(t, 0) for t in raw_type_counts)
    print(f"Indexed {len(existing)} unique resources"
          f"{f' ({dups} duplicate ids collapsed)' if dups else ''}.", flush=True)

    # ---- Write resource counts (total published per type, from the source JSON) ----
    with open(f"resource_counts_{contract}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["resource_type", "unique_count", "raw_count"])
        for t in sorted(type_counts):
            w.writerow([t, type_counts[t], raw_type_counts.get(t, type_counts[t])])

    # ---- Pass 2: check every reference ----
    print("Pass 2: checking references ...", flush=True)
    dangling = []                 # source_type, source_id, file, field, target_type, target_id
    affected = {}                 # (src_type, field, tgt_type) -> source ids with >=1 broken ref
    ref_totals = {}               # (src_type, field, tgt_type) -> total refs seen
    src_with_any_ref = {}         # (src_type, field, tgt_type) -> source ids having such a ref
    bad_ref_counts = {}           # (src_type, field, tgt_type) -> broken ref count
    placeholder_refs = []         # src_type, src_id, file, field, target_type, target_id, connected(bool)
    placeholder_connected = {}    # (src_type, field, tgt_type) -> count where placeholder id resolved
    placeholder_total = {}        # (src_type, field, tgt_type) -> count of placeholder-looking ids seen
    referenced_ids = {}           # (src_type, field, tgt_type) -> set of target ids referenced (any status)

    for path, fname, category, part in files:
        raw = read_bytes(path)
        for r in resources_of(raw):
            src_type, src_id = r.get("resourceType", ""), str(r.get("id", ""))
            for field, ref in find_refs(r):
                parsed = parse_ref(ref)
                if not parsed:
                    continue
                tgt_type, tgt_id = parsed
                key = (src_type, field_root(field), tgt_type)
                ref_totals[key] = ref_totals.get(key, 0) + 1
                src_with_any_ref.setdefault(key, set()).add(src_id)
                referenced_ids.setdefault(key, set()).add(tgt_id)
                connected = (tgt_type, tgt_id) in existing
                if not connected:
                    dangling.append([src_type, src_id, fname, field, tgt_type, tgt_id, ref])
                    affected.setdefault(key, set()).add(src_id)
                    bad_ref_counts[key] = bad_ref_counts.get(key, 0) + 1
                if is_placeholder_id(tgt_id):
                    placeholder_refs.append([src_type, src_id, fname, field, tgt_type, tgt_id,
                                              "yes" if connected else "no"])
                    placeholder_total[key] = placeholder_total.get(key, 0) + 1
                    if connected:
                        placeholder_connected[key] = placeholder_connected.get(key, 0) + 1

    # ---- Write per-reference detail ----
    with open(f"dangling_refs_{contract}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["source_type", "source_id", "source_file", "field_path",
                    "target_type", "target_id", "raw_reference"])
        w.writerows(dangling)

    # ---- Write placeholder-id detail (dummy-looking ids, connected or not) ----
    with open(f"placeholder_refs_{contract}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["source_type", "source_id", "source_file", "field_path",
                    "target_type", "target_id", "connected"])
        w.writerows(placeholder_refs)

    # ---- Write summary (per source.field -> target) ----
    rows = []
    for key in sorted(set(list(ref_totals) + list(affected))):
        src_type, field, tgt_type = key
        n_bad_refs = bad_ref_counts.get(key, 0)
        n_affected = len(affected.get(key, ()))
        n_src_total = type_counts.get(src_type, 0)
        n_src_with_ref = len(src_with_any_ref.get(key, ()))
        pct_of_all = (100.0 * n_affected / n_src_total) if n_src_total else 0.0
        pct_of_linked = (100.0 * n_affected / n_src_with_ref) if n_src_with_ref else 0.0
        # "yes" = target type is published but specific ids are missing;
        # "no"  = that whole resource type was never published in this contract.
        tgt_published = "yes" if tgt_type in published_types else "no"
        rows.append([src_type, field, tgt_type, tgt_published, ref_totals.get(key, 0),
                     n_bad_refs, n_affected, n_src_total, f"{pct_of_all:.2f}",
                     n_src_with_ref, f"{pct_of_linked:.2f}"])

    with open(f"dangling_summary_{contract}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["source_type", "field", "target_type", "target_type_published",
                    "total_refs", "dangling_refs",
                    "affected_source_resources", "total_source_resources", "pct_of_all_source",
                    "source_resources_with_this_ref", "pct_of_linked_source"])
        w.writerows(rows)

    # ---- Console report ----
    print(f"\n--- {contract} dangling references (by field) ---")
    hdr = (f"{'SOURCE.FIELD':<42} {'TARGET':<18} {'PUB':>4} {'REFS':>10} "
           f"{'BROKEN':>10} {'AFFECTED':>9} {'%TYPE':>7}")
    print(hdr)
    print("-" * len(hdr))
    for (src_type, field, tgt_type, tgt_pub, total, bad,
         aff, src_total, pct, _, _) in rows:
        if bad:
            print(f"{src_type + '.' + field:<42} {tgt_type:<18} {tgt_pub:>4} {total:>10,} "
                  f"{bad:>10,} {aff:>9,} {pct:>6}%")
    clean = [r for r in rows if not r[5]]
    if clean:
        print("\nClean (no dangling refs):")
        for src_type, field, tgt_type, *_ in clean:
            print(f"  {src_type}.{field} -> {tgt_type}: all resolve")
    missing_types = sorted({r[2] for r in rows if r[3] == "no" and r[5]})
    if missing_types:
        print(f"\nTarget types referenced but NEVER published: {', '.join(missing_types)}")

    if placeholder_refs:
        print(f"\n--- {contract} placeholder-looking target ids ---")
        phdr = f"{'SOURCE.FIELD':<42} {'TARGET':<18} {'PLACEHOLDER':>11} {'CONNECTED':>10}"
        print(phdr)
        print("-" * len(phdr))
        for key in sorted(set(list(placeholder_total))):
            src_type, field, tgt_type = key
            tot = placeholder_total.get(key, 0)
            conn = placeholder_connected.get(key, 0)
            print(f"{src_type + '.' + field:<42} {tgt_type:<18} {tot:>11,} {conn:>10,}")

        # Roll up by target resource type, across every field that references it.
        by_type_total, by_type_conn = {}, {}
        for row in placeholder_refs:
            _, _, _, _, tgt_type, _, connected = row
            by_type_total[tgt_type] = by_type_total.get(tgt_type, 0) + 1
            if connected == "yes":
                by_type_conn[tgt_type] = by_type_conn.get(tgt_type, 0) + 1
        print(f"\nPlaceholder counts by resource type (all fields combined):")
        thdr = f"{'RESOURCE TYPE':<20} {'PLACEHOLDER':>11} {'CONNECTED':>10} {'DANGLING':>9}"
        print(thdr)
        print("-" * len(thdr))
        for tgt_type in sorted(by_type_total):
            tot = by_type_total[tgt_type]
            conn = by_type_conn.get(tgt_type, 0)
            print(f"{tgt_type:<20} {tot:>11,} {conn:>10,} {tot - conn:>9,}")

        n_conn = sum(placeholder_connected.values())
        n_not_conn = len(placeholder_refs) - n_conn
        print(f"\nTotal placeholder-looking ids: {len(placeholder_refs):,} "
              f"({n_conn:,} still resolve to a real resource, "
              f"{n_not_conn:,} are dangling)")

        if n_not_conn:
            print(f"\nNOT CONNECTED placeholder ids (dangling): {n_not_conn:,}")
            not_conn_by_type = {}
            for row in placeholder_refs:
                _, _, _, _, tgt_type, _, connected = row
                if connected == "no":
                    not_conn_by_type[tgt_type] = not_conn_by_type.get(tgt_type, 0) + 1
            for tgt_type in sorted(not_conn_by_type):
                print(f"  {tgt_type}: {not_conn_by_type[tgt_type]:,}")

        # Placeholder count per source JSON file, standalone (no connected/dangling split).
        by_file = {}
        for row in placeholder_refs:
            fname = row[2]
            by_file[fname] = by_file.get(fname, 0) + 1
        print(f"\nPlaceholder count by source file:")
        for fname in sorted(by_file):
            print(f"  {fname}: {by_file[fname]:,}")

        # End-to-end: InsurancePlan -> Organization specifically (its own chain,
        # not mixed in with every other source type that also references Organization).
        ip_org = [r for r in placeholder_refs
                  if r[0] == "InsurancePlan" and r[4] == "Organization"]
        if ip_org:
            ip_conn = sum(1 for r in ip_org if r[6] == "yes")
            ip_not_conn = len(ip_org) - ip_conn
            print(f"\nEnd-to-end: InsurancePlan -> Organization placeholder ids: "
                  f"{len(ip_org):,} total ({ip_conn:,} connected, {ip_not_conn:,} not connected)")
            for src_type, src_id, fname, field, tgt_type, tgt_id, connected in ip_org:
                print(f"  InsurancePlan/{src_id} --{field}--> Organization/{tgt_id} "
                      f"[{('connected' if connected == 'yes' else 'NOT CONNECTED')}]")

        # End-to-end: OrganizationAffiliation -> Organization specifically.
        oa_org = [r for r in placeholder_refs
                  if r[0] == "OrganizationAffiliation" and r[4] == "Organization"]
        if oa_org:
            oa_conn = sum(1 for r in oa_org if r[6] == "yes")
            oa_not_conn = len(oa_org) - oa_conn
            print(f"\nEnd-to-end: OrganizationAffiliation -> Organization placeholder ids: "
                  f"{len(oa_org):,} total ({oa_conn:,} connected, {oa_not_conn:,} not connected)")
            for src_type, src_id, fname, field, tgt_type, tgt_id, connected in oa_org:
                print(f"  OrganizationAffiliation/{src_id} --{field}--> Organization/{tgt_id} "
                      f"[{('connected' if connected == 'yes' else 'NOT CONNECTED')}]")

        # End-to-end: PractitionerRole -> Practitioner specifically.
        pr_prac = [r for r in placeholder_refs
                   if r[0] == "PractitionerRole" and r[4] == "Practitioner"]
        if pr_prac:
            pr_conn = sum(1 for r in pr_prac if r[6] == "yes")
            pr_not_conn = len(pr_prac) - pr_conn
            print(f"\nEnd-to-end: PractitionerRole -> Practitioner placeholder ids: "
                  f"{len(pr_prac):,} total ({pr_conn:,} connected, {pr_not_conn:,} not connected)")
            for src_type, src_id, fname, field, tgt_type, tgt_id, connected in pr_prac:
                print(f"  PractitionerRole/{src_id} --{field}--> Practitioner/{tgt_id} "
                      f"[{('connected' if connected == 'yes' else 'NOT CONNECTED')}]")

    # ---- Reverse/orphan checks: published resources never referenced back ----
    # e.g. an Organization with no OrganizationAffiliation.organization pointing to it,
    # or a Practitioner with no PractitionerRole.practitioner pointing to it.
    ORPHAN_CHECKS = [
        ("Organization", "OrganizationAffiliation", "organization"),
        ("Practitioner", "PractitionerRole", "practitioner"),
    ]
    orphan_rows = []
    for tgt_type, ref_src_type, field in ORPHAN_CHECKS:
        key = (ref_src_type, field, tgt_type)
        referenced = referenced_ids.get(key, set())
        all_ids = {rid for (rt, rid) in existing if rt == tgt_type}
        orphans = sorted(all_ids - referenced)
        if not all_ids or not orphans:
            continue
        print(f"\n--- {contract}: {tgt_type} not referenced by any {ref_src_type}.{field} ---")
        print(f"{len(orphans):,} of {len(all_ids):,} {tgt_type} resources are orphaned "
              f"({100.0 * len(orphans) / len(all_ids):.2f}%)")
        for oid in orphans:
            meta = org_meta.get((tgt_type, oid), {})
            otype = meta.get("type", "")
            flag = "Expected" if is_expected_orphan_type(otype) else "Review"
            orphan_rows.append([tgt_type, oid, identifiers.get((tgt_type, oid), ""),
                                 meta.get("name", ""), otype, flag,
                                 ref_src_type, field])

    if orphan_rows:
        with open(f"orphan_refs_{contract}.csv", "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["orphan_type", "orphan_id", "identifier", "name", "org_type", "flag",
                        "expected_referencing_type", "expected_field"])
            w.writerows(orphan_rows)

    # ---- Console: the two headline connectivity checks only ----
    org_orphan_n = sum(1 for r in orphan_rows
                        if r[0] == "Organization" and r[6] == "OrganizationAffiliation"
                        and r[7] == "organization")
    prac_orphan_n = sum(1 for r in orphan_rows
                         if r[0] == "Practitioner" and r[6] == "PractitionerRole"
                         and r[7] == "practitioner")
    print(f"\nOrphan Connectivity Checks")
    print(f"Organization is connected to OrganizationAffiliation"
          + (f" - ({org_orphan_n:,} orphaned)" if org_orphan_n else ""))
    print(f"Practitioner is connected to PractitionerRole"
          + (f" - ({prac_orphan_n:,} orphaned)" if prac_orphan_n else ""))

    print(f"\nTotal dangling references: {len(dangling):,}")
    print(f"Wrote: dangling_refs_{contract}.csv, dangling_summary_{contract}.csv"
          + (f", placeholder_refs_{contract}.csv" if placeholder_refs else "")
          + (f", orphan_refs_{contract}.csv" if orphan_rows else ""))
    return len(dangling)


def _read_csv(path):
    if not os.path.exists(path):
        return [], []
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    return (rows[0], rows[1:]) if rows else ([], [])


def build_report(contracts):
    """After each contract's own run has written its CSVs, roll them into one
    manager-facing report: an Excel workbook (per-contract detail tabs) and a
    Word summary. Each contract's numbers still come from its own separate run;
    this just collates the already-written per-contract outputs."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from docx import Document
    except ImportError as e:
        print(f"\nSkipping report generation (missing dependency: {e}). "
              f"Run: pip install openpyxl python-docx")
        return

    # Dynamic filename driven by whichever contract(s) were actually run, e.g.
    # reference_integrity_contract_H5826_report.xlsx for one contract, or
    # reference_integrity_contract_H1619_H3124_report.xlsx for several.
    report_base = "reference_integrity_contract_" + "_".join(contracts) + "_report"

    data = {}
    for c in contracts:
        shdr, srows = _read_csv(f"dangling_summary_{c}.csv")
        ohdr, orows = _read_csv(f"orphan_refs_{c}.csv")
        chdr, crows = _read_csv(f"resource_counts_{c}.csv")
        phdr, prows = _read_csv(f"placeholder_refs_{c}.csv")
        total_dangling = sum(int(r[5]) for r in srows) if srows else 0
        total_affected = sum(int(r[6]) for r in srows) if srows else 0
        # placeholder rollup by target resource type: total seen, connected, dangling
        ph_by_type = {}
        for row in prows:
            _, _, _, _, tgt_type, _, connected = row
            e = ph_by_type.setdefault(tgt_type, {"total": 0, "connected": 0})
            e["total"] += 1
            if connected == "yes":
                e["connected"] += 1
        # Fully dynamic: group orphan rows by the actual (orphan_type -> referencing
        # relationship) they belong to, whatever those turn out to be -- not a
        # hardcoded Organization/Practitioner pair. Column 4 (org_type / e.g.
        # Payer, Network) explains *why* each group is orphaned.
        conn_checks = {}
        for r in orows:
            orphan_type, _, _, _, otype, flag, ref_type, field = r[:8]
            key = (orphan_type, ref_type, field)
            e = conn_checks.setdefault(key, {"count": 0, "types": {}, "review": []})
            e["count"] += 1
            if otype:
                e["types"][otype] = e["types"].get(otype, 0) + 1
            if flag == "Review":
                e["review"].append(r)
        data[c] = dict(shdr=shdr, srows=srows, ohdr=ohdr, orows=orows,
                        crows=crows, ph_by_type=ph_by_type,
                        total_dangling=total_dangling, total_affected=total_affected,
                        conn_checks=conn_checks)

    grand_dangling = sum(d["total_dangling"] for d in data.values())
    grand_orphans = sum(len(d["orows"]) for d in data.values())

    # ---- Excel ----
    wb = Workbook()
    HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
    HEADER_FONT = Font(color="FFFFFF", bold=True)
    BAD_FILL = PatternFill("solid", fgColor="FCE4E4")
    OK_FILL = PatternFill("solid", fgColor="E4F7E4")

    def style_header(ws, row=1, ncols=1):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def autosize(ws, ncols):
        for col in range(1, ncols + 1):
            letter = get_column_letter(col)
            maxlen = max((len(str(c.value)) if c.value is not None else 0)
                         for c in ws[letter])
            ws.column_dimensions[letter].width = min(max(maxlen + 2, 10), 60)

    ws = wb.active
    ws.title = "Overview"
    ws.append(["FHIR Provider Directory - Reference Integrity Report"])
    ws["A1"].font = Font(size=14, bold=True)
    ws.append([f"Generated: {date.today().isoformat()}"])
    ws.append([])
    ws.append(["Contract", "Dangling References", "Orphaned Organizations",
               "Total Referenced Resources", "Status"])
    style_header(ws, row=4, ncols=5)
    for c in contracts:
        d = data[c]
        total_refs = sum(int(r[4]) for r in d["srows"]) if d["srows"] else 0
        status = "CLEAN" if d["total_dangling"] == 0 else "ISSUES FOUND"
        ws.append([c, d["total_dangling"], len(d["orows"]), total_refs, status])
        r = ws.max_row
        fill = OK_FILL if d["total_dangling"] == 0 else BAD_FILL
        for col in range(1, 6):
            ws.cell(row=r, column=col).fill = fill
    ws.append([])
    ws.append(["TOTAL (all contracts)", grand_dangling, grand_orphans])
    ws[f"A{ws.max_row}"].font = Font(bold=True)
    autosize(ws, 5)

    for c in contracts:
        d = data[c]
        wsum = wb.create_sheet(f"{c} Summary")
        if d["shdr"]:
            wsum.append(d["shdr"])
            style_header(wsum, row=1, ncols=len(d["shdr"]))
            for row in d["srows"]:
                wsum.append(row)
                rr = wsum.max_row
                if int(row[5]) > 0:
                    for col in range(1, len(d["shdr"]) + 1):
                        wsum.cell(row=rr, column=col).fill = BAD_FILL
            autosize(wsum, len(d["shdr"]))
        wsum.freeze_panes = "A2"

        worp = wb.create_sheet(f"{c} Orphans")
        if d["ohdr"]:
            worp.append(d["ohdr"])
            style_header(worp, row=1, ncols=len(d["ohdr"]))
            for row in d["orows"]:
                worp.append(row)
            autosize(worp, len(d["ohdr"]))
        else:
            worp.append(["No orphaned resources found."])
        worp.freeze_panes = "A2"

    xlsx_path = f"{report_base}.xlsx"
    wb.save(xlsx_path)

    # ---- Word ----
    def add_table(doc, headers, rows, style="Light Grid Accent 1"):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = style
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = h
            cell.paragraphs[0].runs[0].font.bold = True
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
        return t

    today_str = date.today().strftime("%d-%b-%Y")

    doc = Document()
    doc.add_heading("Reference Integrity Validation Report", level=0)
    doc.add_paragraph(f"Execution Date: {today_str}")
    doc.add_paragraph("Validation Type: Forward Reference / Dangling Reference Validation")
    p = doc.add_paragraph()
    p.add_run("Contracts covered: ").bold = True
    p.add_run(", ".join(contracts))

    overall_status = "PASS" if grand_dangling == 0 else "FAIL"
    doc.add_heading("1. Overall Summary", level=1)
    add_table(doc, ["Metric", "Result"], [
        ["Execution Date", today_str],
        ["Environment", "Production"],
        ["Validation Type", "Forward Reference / Dangling Reference"],
        ["Contracts Validated", len(contracts)],
        ["Dangling References", f"{grand_dangling:,}"],
        ["Affected Source Resources",
         f"{sum(d['total_affected'] for d in data.values()):,}"],
        ["Orphaned Organizations", f"{grand_orphans:,}"],
        ["Reference Integrity", overall_status],
    ])

    for c in contracts:
        d = data[c]
        status = "PASS" if d["total_dangling"] == 0 else "FAIL"
        n_relationships = len(d["srows"])

        doc.add_heading(f"Contract: {c}", level=1)

        doc.add_heading("1. Execution Summary", level=2)
        add_table(doc, ["Metric", "Result"], [
            ["Execution Date", today_str],
            ["Environment", "Production"],
            ["Validation Type", "Forward Reference / Dangling Reference"],
            ["Relationship Checks", n_relationships],
            ["Dangling References", f"{d['total_dangling']:,}"],
            ["Affected Source Resources", f"{d['total_affected']:,}"],
            ["Reference Integrity", status],
        ])

        doc.add_heading("2. Validation Results", level=2)
        vrows = []
        for row in d["srows"]:
            src_type, field, tgt_type, tgt_pub, total, bad, aff = row[:7]
            vrows.append([src_type, field, tgt_type, f"{int(total):,}",
                          bad, aff, "PASS" if int(bad) == 0 else "FAIL"])
        add_table(doc, ["Source Type", "Field", "Target Type", "Total References",
                        "Dangling", "Affected Resources", "Result"], vrows)

        doc.add_heading("3. Resource Counts (Total in JSON)", level=2)
        if d["crows"]:
            add_table(doc, ["Resource Type", "Unique Count", "Raw Count"], d["crows"])
        else:
            doc.add_paragraph("No resource-count data available for this contract.")

        doc.add_heading("4. Placeholder-Looking Target IDs", level=2)
        if d["ph_by_type"]:
            ph_rows = []
            for tgt_type in sorted(d["ph_by_type"]):
                e = d["ph_by_type"][tgt_type]
                dangling_ph = e["total"] - e["connected"]
                ph_rows.append([tgt_type, e["total"], e["connected"], dangling_ph])
            add_table(doc, ["Target Type", "Placeholder-Looking", "Connected", "Dangling"],
                      ph_rows)
        else:
            doc.add_paragraph("No placeholder-looking ids found.")

        doc.add_heading("5. Orphan Connectivity Checks", level=2)
        # Fully dynamic: one line per (orphan_type -> referencing relationship)
        # actually observed for this contract. An orphan of an EXPECTED_ORPHAN_TYPES
        # type (e.g. Payer, Network) stays Pass -- it's not meant to be linked.
        # Anything else is flagged REVIEW so a real gap doesn't hide behind a Pass.
        any_review = False
        if d["conn_checks"]:
            for (orphan_type, ref_type, field), e in sorted(d["conn_checks"].items()):
                type_note = ""
                if e["types"]:
                    breakdown = ", ".join(f"{t}: {n}" for t, n in sorted(e["types"].items()))
                    type_note = f" -- type(s): {breakdown}"
                if e["review"]:
                    any_review = True
                    status = f"REVIEW ({len(e['review'])} of {e['count']} not an expected type)"
                else:
                    status = "Pass"
                doc.add_paragraph(
                    f"{orphan_type} is connected to {ref_type}.{field} - "
                    f"{status} ({e['count']} orphaned{type_note})"
                )
                if e["review"]:
                    rt = doc.add_table(rows=1, cols=4)
                    rt.style = "Light List Accent 2"
                    rh = rt.rows[0].cells
                    for i, h in enumerate(["Orphan ID", "Identifier", "Name", "Type"]):
                        rh[i].text = h
                        rh[i].paragraphs[0].runs[0].font.bold = True
                    for row in e["review"]:
                        rc = rt.add_row().cells
                        rc[0].text = row[1]
                        rc[1].text = row[2]
                        rc[2].text = row[3]
                        rc[3].text = row[4] or "(no type)"
            doc.add_paragraph(
                "Note: orphans of an expected type (Payer, Network) are normal -- "
                "those Organization levels are not required to be linked via an "
                "OrganizationAffiliation record. Any REVIEW line above lists "
                "orphans of an unexpected type that likely need a data fix. "
                "See the Orphan Details table below for the full breakdown."
            )
        else:
            doc.add_paragraph("Every published resource is referenced back at least "
                               "once by its expected relationship - Pass.")

        doc.add_heading("6. Orphan Details", level=2)
        ORPHAN_TABLE_CAP = 500
        if d["orows"]:
            shown = d["orows"][:ORPHAN_TABLE_CAP]
            add_table(doc, ["Orphan Type", "Orphan ID", "Identifier", "Name",
                            "Org/Practitioner Type", "Flag",
                            "Expected Referencing Type", "Expected Field"], shown)
            if len(d["orows"]) > ORPHAN_TABLE_CAP:
                doc.add_paragraph(
                    f"... {len(d['orows']) - ORPHAN_TABLE_CAP:,} more not shown here -- "
                    f"see orphan_refs_{c}.csv for the full list."
                )
        else:
            doc.add_paragraph("No orphaned resources found for this contract.")

        doc.add_paragraph()

    doc.add_heading("Recommendation", level=1)
    doc.add_paragraph(
        "Review the orphaned records above with the data provider to confirm "
        "whether they are intentionally unaffiliated or should be linked via an "
        "OrganizationAffiliation / PractitionerRole entry." if grand_orphans else
        "No follow-up required -- all references resolve and no orphaned "
        "resources were found."
    )
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.add_run("Full detail (per-reference breakdowns) is available in the "
                 f"accompanying Excel workbook: {xlsx_path}").italic = True

    docx_path = f"{report_base}.docx"
    doc.save(docx_path)

    print(f"\nWrote report: {xlsx_path}, {docx_path}")




CONTRACT_DIR_RE = re.compile(r"^[hH]\d+$")


def discover_contracts(base_dir):
    """Every subfolder of base_dir named like a contract id (H1619, H5826, ...)
    that actually holds at least one *.json bundle, mapped to its full path."""
    contracts = {}
    for name in sorted(os.listdir(base_dir)):
        full = os.path.join(base_dir, name)
        if os.path.isdir(full) and CONTRACT_DIR_RE.match(name):
            if glob.glob(os.path.join(full, "*.json")):
                contracts[name.upper()] = full
    return contracts


# CLI: python check_hy.py [CONTRACT]
#   CONTRACT  run just one contract, e.g. H1619 (default: all contracts found
#             as sibling folders of this script, e.g. ./H1619/*.json)
SCRIPT_DIR = os.environ.get("INPUT_DIR") or os.path.dirname(os.path.abspath(__file__))  # INPUT_DIR: set by GitHub Actions runner
CONTRACTS = discover_contracts(SCRIPT_DIR)

args = [a for a in sys.argv[1:]]
positional = [a for a in args if not a.startswith("-")]
arg = positional[0].upper() if positional else None

if not CONTRACTS:
    print(f"No contract folders with *.json files found under {SCRIPT_DIR}.")
    print("Expected e.g. .\\H1619\\*.json next to this script.")
    sys.exit(1)

if arg:
    if arg not in CONTRACTS:
        print(f"Unknown/empty contract '{arg}'. Choose from: {', '.join(CONTRACTS)}")
        sys.exit(1)
    targets = {arg: CONTRACTS[arg]}
else:
    targets = CONTRACTS

grand_total = 0
for c, folder in targets.items():
    grand_total += run(c, folder)
if len(targets) > 1:
    print(f"\n==== ALL CONTRACTS: {grand_total:,} total dangling references ====")
build_report(list(targets))   # Excel + Word report (ported from maplan check_refs.py)
sys.exit(1 if grand_total else 0)  # non-zero exit = dangling references found (GitHub Actions FAIL)
